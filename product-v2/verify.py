"""Independent check of the spreadsheet's dates and amounts.
Builds scenarios, lets LibreOffice calculate the real workbook, then compares every value
against a plain-Python model written separately from the spreadsheet formulas."""
import calendar, datetime as dt, math, os, shutil, subprocess, json, sys
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "The Paycheck Budget System - Spreadsheet.xlsx")
RECALC = "/mnt/skills/public/xlsx/scripts/recalc.py"
D = dt.date

def clamp(y, m, d):
    while m > 12: y, m = y + 1, m - 12
    while m < 1: y, m = y - 1, m + 12
    return D(y, m, min(d, calendar.monthrange(y, m)[1]))

def add_months(base, k):
    return clamp(base.year, base.month + k, base.day)

# ---------- independent model ----------
def paydays(anchor, freq, count_extra=2):
    end = add_months(anchor, 12)
    out = []
    if freq in ("Weekly", "Every 2 weeks"):
        step = 7 if freq == "Weekly" else 14
        d = anchor
        while len(out) < 60:
            out.append(d); d += dt.timedelta(days=step)
    elif freq == "Monthly":
        out = [add_months(anchor, k) for k in range(60)]
    else:  # twice a month: two fixed days each month
        a = anchor.day
        days = sorted({a, a + 15 if a <= 15 else a - 15})
        y, m = anchor.year, anchor.month
        while len(out) < 60:
            for dd in days:
                x = clamp(y, m, dd)
                if x >= anchor and x not in out:
                    out.append(x)
            m += 1
            if m > 12: y, m = y + 1, 1
        out.sort()
    inside = [d for d in out if d < end]
    nxt = out[len(inside)] if len(out) > len(inside) else None
    return inside, out

def occurrences(due, freq, upto):
    res = []
    if freq in ("Weekly", "Every 2 weeks"):
        step = 7 if freq == "Weekly" else 14
        d = due
        while d < upto:
            res.append(d); d += dt.timedelta(days=step)
    else:
        mstep = {"Monthly": 1, "Quarterly": 3, "Twice a year": 6, "Yearly": 12}[freq]
        k = 0
        while True:
            d = add_months(due, k * mstep)
            if d >= upto: break
            res.append(d); k += 1
    return res

PER_YEAR = {"Weekly": 52, "Every 2 weeks": 26, "Monthly": 12, "Quarterly": 4, "Twice a year": 2, "Yearly": 1}
PAY_PER_YEAR = {"Weekly": 52, "Every 2 weeks": 26, "Twice a month": 24, "Monthly": 12}
DIVIDE = {"Weekly": 48, "Every 2 weeks": 24, "Twice a month": 24, "Monthly": 12}
USUAL = {"Weekly": 4, "Every 2 weeks": 2, "Twice a month": 2, "Monthly": 1}

def model(sc):
    inside, allp = paydays(sc["anchor"], sc["freq"])
    periods = []
    for i, p in enumerate(inside):
        end = allp[i + 1]
        periods.append((p, end))
    grid = []
    for (s, e) in periods:
        row = []
        for b in sc["bills"]:
            name, freq, amt, due = b
            if due is None:
                row.append(0); continue
            n = sum(1 for o in occurrences(due, freq, e) if s <= o < e)
            row.append(n * amt)
        grid.append(row)
    per = []
    for name, freq, amt, due in sc["bills"]:
        div = PAY_PER_YEAR[sc["freq"]] if freq in ("Weekly", "Every 2 weeks") else DIVIDE[sc["freq"]]
        per.append(math.ceil(amt * PER_YEAR[freq] / div - 1e-9))
    ready = []
    for (name, freq, amt, due), pp in zip(sc["bills"], per):
        if due is None: ready.append(None); continue
        cnt = sum(1 for p in inside if p <= due)
        ready.append(max(0, amt - pp * cnt))
    extra = []
    for p in inside:
        idx = sum(1 for q in inside if q.year == p.year and q.month == p.month and q <= p)
        extra.append(idx > USUAL[sc["freq"]])
    return dict(paydays=inside, grid=grid, per=per, ready=ready, extra=extra)

# ---------- run the real workbook ----------
def todate(v):
    if v in (None, ""): return None
    if isinstance(v, dt.datetime): return v.date()
    if isinstance(v, (int, float)): return D(1899, 12, 30) + dt.timedelta(days=int(v))
    return v

def sheet_values(sc, path):
    wb = load_workbook(SRC)
    m = wb["My Budget"]
    m["B6"] = sc["freq"]; m["B7"] = sc["anchor"]
    for r in range(14, 44):
        for c in range(1, 6): m.cell(row=r, column=c).value = None
    for i, (name, freq, amt, due) in enumerate(sc["bills"]):
        r = 14 + i
        m.cell(row=r, column=1, value=name); m.cell(row=r, column=2, value="Bills")
        m.cell(row=r, column=3, value=freq); m.cell(row=r, column=4, value=amt); m.cell(row=r, column=5, value=due)
    wb.save(path)
    out = subprocess.run([sys.executable, RECALC, path, "180"], capture_output=True, text=True).stdout
    res = json.loads(out)
    assert res.get("status") == "success", res
    v = load_workbook(path, data_only=True)
    m, pc = v["My Budget"], v["Pay Calendar"]
    pays, grid, extra = [], [], []
    for r in range(8, 62):
        d = todate(pc.cell(row=r, column=2).value)
        if d is None: continue
        pays.append(d)
        extra.append(pc.cell(row=r, column=3).value == "★")
        grid.append([pc.cell(row=r, column=9 + j).value or 0 for j in range(len(sc["bills"]))])
    per = [m.cell(row=14 + i, column=6).value for i in range(len(sc["bills"]))]
    ready = [m.cell(row=14 + i, column=7).value for i in range(len(sc["bills"]))]
    ready = [None if x in (None, "") else x for x in ready]
    return dict(paydays=pays, grid=grid, per=per, ready=ready, extra=extra)

BILLS = [
    ("Monthly on the 31st", "Monthly", 100, D(2027, 1, 31)),
    ("Monthly on the 29th", "Monthly", 60, D(2027, 11, 29)),
    ("Monthly on the 1st", "Monthly", 1800, D(2026, 11, 1)),
    ("Quarterly on 30 Nov", "Quarterly", 150, D(2026, 11, 30)),
    ("Twice a year, 29 Feb", "Twice a year", 300, D(2028, 2, 29)),
    ("Yearly, 29 Feb 2028", "Yearly", 400, D(2028, 2, 29)),
    ("Yearly, next month", "Yearly", 720, D(2026, 12, 10)),
    ("Yearly, past date", "Yearly", 250, D(2025, 3, 15)),
    ("Weekly", "Weekly", 30, D(2026, 10, 19)),
    ("Every 2 weeks", "Every 2 weeks", 85, D(2026, 10, 23)),
    ("Weekly from 2028", "Weekly", 20, D(2028, 2, 27)),
    ("No due date", "Monthly", 160, None),
    ("Far future", "Yearly", 500, D(2031, 6, 1)),
]
SCEN = [
    ("Every 2 weeks, from 16 Oct 2026", "Every 2 weeks", D(2026, 10, 16)),
    ("Weekly, from 31 Dec 2027 (into leap year)", "Weekly", D(2027, 12, 31)),
    ("Monthly, on the 31st, from 31 Jan 2028", "Monthly", D(2028, 1, 31)),
    ("Monthly, on the 29th, from 29 Feb 2028", "Monthly", D(2028, 2, 29)),
    ("Twice a month, 5th & 20th", "Twice a month", D(2026, 10, 20)),
    ("Twice a month, 1st & 16th", "Twice a month", D(2026, 11, 1)),
    ("Twice a month, 15th & 30th (Feb)", "Twice a month", D(2027, 1, 15)),
    ("Twice a month, 16th & 31st", "Twice a month", D(2027, 1, 31)),
    ("Every 2 weeks, from 29 Feb 2028", "Every 2 weeks", D(2028, 2, 29)),
    ("Weekly, from 16 Oct 2026", "Weekly", D(2026, 10, 16)),
]
os.makedirs(os.path.join(HERE, "verify"), exist_ok=True)
total_checks, fails = 0, []
for name, freq, anchor in SCEN:
    sc = dict(freq=freq, anchor=anchor, bills=BILLS)
    exp = model(sc)
    got = sheet_values(sc, os.path.join(HERE, "verify", "case.xlsx"))
    checks = 0
    def chk(label, a, b):
        global total_checks
        nonlocal_c[0] += 1
        if a != b: fails.append((name, label, a, b))
    nonlocal_c = [0]
    chk("number of paydays", len(exp["paydays"]), len(got["paydays"]))
    for i, (a, b) in enumerate(zip(exp["paydays"], got["paydays"])): chk(f"payday {i+1}", a, b)
    for i, (a, b) in enumerate(zip(exp["extra"], got["extra"])): chk(f"extra flag {i+1}", a, b)
    for i, (ra, rb) in enumerate(zip(exp["grid"], got["grid"])):
        for j, (a, b) in enumerate(zip(ra, rb)): chk(f"pay {i+1} / {BILLS[j][0]}", a, b)
    for j, (a, b) in enumerate(zip(exp["per"], got["per"])): chk(f"per payday / {BILLS[j][0]}", a, b)
    for j, (a, b) in enumerate(zip(exp["ready"], got["ready"])): chk(f"top-up / {BILLS[j][0]}", a, b)
    total_checks += nonlocal_c[0]
    print(f"{name}: {len(got['paydays'])} paydays, {sum(got['extra'])} extra, {nonlocal_c[0]} values checked")
print("TOTAL CHECKED", total_checks, "MISMATCHES", len(fails))
for f in fails[:40]: print("  ", f)
