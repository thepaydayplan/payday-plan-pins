import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.chart import LineChart, Reference

OUT = "The Paycheck Budget System - Spreadsheet.xlsx"
NAVY, GREEN, GOLD, RED, BROWN, BLACK = "1B3A6B", "1E7A4C", "C9A227", "B23A3A", "7A5A0F", "121212"
CREAM, OFF, INPUT = "F3EFE4", "FAF8F3", "FFF4CC"
TINT = {"Bills": "E6ECF5", "Debt": "ECEAE4", "Savings": "E2F0E7", "Don't Touch": "F7E6E6", "Everyday": "F6EED3", "Buffer": "FBF3D9"}
BCOL = {"Bills": NAVY, "Debt": BLACK, "Savings": GREEN, "Don't Touch": RED, "Everyday": BROWN, "Buffer": "8A6D12"}
F = "Arial"
MONEY = '$#,##0;-$#,##0;"-"'
thin = Side(style="thin", color="D9D3C4")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

def font(**k):
    k.setdefault("name", F); k.setdefault("size", 10)
    return Font(**k)

def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)

def title(ws, t, sub, width_cols=8):
    ws["A1"] = t; ws["A1"].font = font(size=20, bold=True, color=NAVY)
    ws["A2"] = sub; ws["A2"].font = font(size=11, color=GREEN, bold=True)
    ws.row_dimensions[1].height = 30
    ws.sheet_view.showGridLines = False

def header(ws, row, labels, col=1, bg=NAVY, fg="FFFFFF"):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=l)
        c.font = font(bold=True, color=fg, size=9.5); c.fill = fill(bg)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28

def section(ws, row, text, color=NAVY):
    c = ws.cell(row=row, column=1, value=text)
    c.font = font(bold=True, size=11, color=color)

def inp(c, fmt=None):
    c.fill = fill(INPUT); c.font = font(color="0000FF"); c.border = BOX
    if fmt: c.number_format = fmt

def calc(c, fmt=None, bold=False, color="000000"):
    c.font = font(bold=bold, color=color); c.border = BOX
    if fmt: c.number_format = fmt

def widths(ws, ws_w):
    for k, v in ws_w.items():
        ws.column_dimensions[k].width = v

wb = Workbook()

# ======================= START HERE =======================
s = wb.active; s.title = "Start Here"
title(s, "The Paycheck Budget System", "Companion spreadsheet · The Payday Plan")
widths(s, {"A": 5, "B": 100})
rows = [
    ("HOW TO USE IT", None),
    ("1", "Go to the Setup tab. Type your pay at the top, then list every bill in the Bill Finder. Example numbers for a family paid $2,500 every two weeks are already filled in, so you can see how it works. Type over them with yours."),
    ("2", "Setup works out what each bucket needs every payday, and your Payday Number: what moves automatically, what's yours to live on, and what stays spare."),
    ("3", "Year of Paydays fills itself in from your next payday. Extra paydays get a star, so you can plan where they go before they land."),
    ("4", "Every payday, open the Payday tab and do the 6 steps in order. Tick each one. Then add a line to the Payday Log, including how many minutes it took. Watch that number drop."),
    ("5", "The 30-Day Challenge, Check-In and Trackers tabs match pages 13 to 15 of the workbook. No reprinting needed."),
    ("", None),
    ("THE COLOURS", None),
    ("", "Yellow cells with blue text are yours to type in."),
    ("", "White cells work themselves out. Please don't type over them."),
    ("", None),
    ("GOOGLE SHEETS", None),
    ("", "Upload this file to Google Drive, then open it with Google Sheets. Everything works the same way."),
    ("", None),
    ("THE LINE TO REMEMBER", None),
    ("", "Pay the bills. Pay the debt. Pay yourself. Guard the rest."),
    ("", None),
    ("", "General information, not financial advice. For personal use only. © The Payday Plan"),
]
r = 4
for a, b in rows:
    if b is None and a:
        s.cell(row=r, column=1, value=a).font = font(bold=True, size=11, color=NAVY)
    else:
        if a:
            c = s.cell(row=r, column=1, value=int(a)); c.font = font(bold=True, color="FFFFFF"); c.fill = fill(NAVY)
            c.alignment = Alignment(horizontal="center", vertical="top")
        if b:
            c = s.cell(row=r, column=2, value=b); c.font = font(size=10.5); c.alignment = Alignment(wrap_text=True, vertical="top")
            if b.startswith("Yellow"):
                c.fill = fill(INPUT); c.font = font(size=10.5, color="0000FF")
            if b.startswith("Pay the bills"):
                c.font = font(size=13, bold=True, color=GREEN)
            if b.startswith("General"):
                c.font = font(size=9, color="6B6657")
            s.row_dimensions[r].height = 30 if len(b) > 100 else 18
    r += 1

# ======================= SETUP =======================
st = wb.create_sheet("Setup")
title(st, "Bucket Setup", "Fill in the yellow cells once. Everything else works itself out.")
widths(st, {"A": 30, "B": 16, "C": 16, "D": 14, "E": 14, "F": 16, "G": 14, "H": 30})
section(st, 4, "YOUR PAY")
st["A5"] = "Paycheck amount"; st["B5"] = 2500; inp(st["B5"], MONEY)
st["A6"] = "How often you're paid"; st["B6"] = "Every 2 weeks"; inp(st["B6"])
st["A7"] = "Your next payday"; st["B7"] = dt.date(2026, 10, 16); inp(st["B7"], "ddd d mmm yyyy")
st["A8"] = "Buffer you want to keep"; st["B8"] = 100; inp(st["B8"], MONEY)
st["A9"] = "Paydays a year"; st["B9"] = '=IF(B6="Weekly",52,IF(B6="Every 2 weeks",26,IF(B6="Twice a month",24,12)))'; calc(st["B9"], "0")
st["A10"] = "Divide-by number"; st["B10"] = '=IF(B6="Weekly",48,IF(B6="Every 2 weeks",24,IF(B6="Twice a month",24,12)))'; calc(st["B10"], "0")
st["A11"] = "Usual paydays a month"; st["B11"] = '=IF(B6="Weekly",4,IF(B6="Monthly",1,2))'; calc(st["B11"], "0")
st["C10"] = "Monthly-and-up bills are divided by 48 or 24 (not 52 or 26), so you save a little extra every year."
st["C10"].font = font(size=9, italic=True, color="6B6657")
for a in range(5, 12):
    st[f"A{a}"].font = font(bold=True)
dv_freq = DataValidation(type="list", formula1='"Weekly,Every 2 weeks,Twice a month,Monthly"', allow_blank=False)
st.add_data_validation(dv_freq); dv_freq.add("B6")

section(st, 13, "YOUR PAYDAY NUMBER")
header(st, 14, ["Bucket", "Per payday", "Account name", "", "Auto transfer set?", "", "", "What it's for"])
BR1, BR2 = 37, 86   # bill rows
buckets = [
    ("1  Bills", f'=SUMIF($B${BR1}:$B${BR2},"Bills",$G${BR1}:$G${BR2})', "CAR, UTILITIES, PHONE", "Every bill gets the paycheck before it's due."),
    ("2  Debt", f'=SUMIF($B${BR1}:$B${BR2},"Debt",$G${BR1}:$G${BR2})', "Main account", "Every minimum on time. Extra goes to the smallest balance."),
    ("3  Savings", f'=SUMIF($B${BR1}:$B${BR2},"Savings",$G${BR1}:$G${BR2})', "SAVINGS", "Paid like a bill. Even $10 counts."),
    ("4  Don't Touch", f'=SUMIF($B${BR1}:$B${BR2},"Don\'t Touch",$G${BR1}:$G${BR2})', "DON'T TOUCH", "Only for things that break. Never for a sale."),
    ("5  Everyday", "=B33", "Main account", "Planned from what's left. Never first."),
    ("6  Buffer", "=B5-SUM(B15:B19)", "Main account", "Stays spare until next payday."),
]
names = ["Bills", "Debt", "Savings", "Don't Touch", "Everyday", "Buffer"]
dv_tick = DataValidation(type="list", formula1='"✓"', allow_blank=True)
st.add_data_validation(dv_tick)
for i, (lab, f_, acct, what) in enumerate(buckets):
    rr = 15 + i
    nm = names[i]
    c = st.cell(row=rr, column=1, value=lab); c.font = font(bold=True, color=BCOL[nm]); c.fill = fill(TINT[nm]); c.border = BOX
    c = st.cell(row=rr, column=2, value=f_); calc(c, MONEY, bold=True); c.fill = fill(TINT[nm])
    c = st.cell(row=rr, column=3, value=acct if i < 4 else None)
    if i < 4:
        inp(c); st.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=4)
        c = st.cell(row=rr, column=5, value="✓" if i < 4 else None); inp(c); c.alignment = Alignment(horizontal="center"); dv_tick.add(c.coordinate)
    c = st.cell(row=rr, column=8, value=what); c.font = font(size=9, color="5B6475")
st["A21"] = '=IF(B20>=B8,"✓ Your buffer is healthy.","⚠ Buffer is under your target. Trim Everyday first, then extra debt payments, then savings (never to zero). Never skip a bill.")'
st["A21"].font = font(bold=True, color=GREEN)
st.merge_cells("A21:H21")
st["A23"] = '="Every payday, "&TEXT(SUM(B15:B18),"$#,##0")&" moves automatically into buckets 1 to 4."'
st["A24"] = '=TEXT(B19,"$#,##0")&" is yours to live on until next payday."'
st["A25"] = '="And "&TEXT(B20,"$#,##0")&" stays spare."'
for a in ("A23", "A24", "A25"):
    st[a].font = font(size=13, bold=True, color=NAVY); st.merge_cells(f"{a}:{a.replace('A', 'H')}")
st.conditional_formatting.add("A21", FormulaRule(formula=["$B$20<$B$8"], font=Font(name=F, bold=True, color=RED)))
st.conditional_formatting.add("B20", CellIsRule(operator="lessThan", formula=["$B$8"], fill=fill("F7E6E6"), font=Font(name=F, bold=True, color=RED)))

section(st, 27, "EVERYDAY PLAN (PER PAYDAY)", BROWN)
for i, (k, v) in enumerate([("Groceries", 450), ("Gas", 150), ("Kids & household", 150), ("Fun money", 120)]):
    st.cell(row=28 + i, column=1, value=k).font = font()
    c = st.cell(row=28 + i, column=2, value=v); inp(c, MONEY)
st["A32"] = "(anything else)"; st["A32"].font = font(color="6B6657", italic=True); inp(st["B32"], MONEY)
st["A33"] = "Everyday total"; st["A33"].font = font(bold=True, color=BROWN); st["B33"] = "=SUM(B28:B32)"; calc(st["B33"], MONEY, bold=True)

section(st, 35, "THE BILL FINDER: EVERY BILL, EVEN THE SNEAKY ONES")
header(st, 36, ["Bill", "Bucket", "How often", "Amount", "Due", "Cost a year", "Per payday", "Notes"])
bills = [
    ("Rent", "Bills", "Monthly", 1800, "1st"), ("Power", "Bills", "Monthly", 160, "20th"),
    ("Phone & internet", "Bills", "Monthly", 140, "5th"), ("Car insurance", "Bills", "Yearly", 1440, "March"),
    ("Car registration", "Bills", "Yearly", 360, "March"), ("Water", "Bills", "Quarterly", 150, "Mar/Jun/Sep/Dec"),
    ("Car service & tires", "Bills", "Yearly", 600, "Any time"), ("Streaming & apps", "Bills", "Monthly", 50, "Various"),
    ("Credit card minimum", "Debt", "Monthly", 200, "12th"),
    ("Savings", "Savings", "Monthly", 160, ""), ("Christmas", "Savings", "Yearly", 720, "December"),
    ("Birthdays & gifts", "Savings", "Yearly", 480, "Various"), ("School costs", "Savings", "Yearly", 480, "Jan & Jul"),
    ("Don't Touch fund", "Don't Touch", "Monthly", 160, ""),
]
dv_b = DataValidation(type="list", formula1='"Bills,Debt,Savings,Don\'t Touch"', allow_blank=True)
dv_o = DataValidation(type="list", formula1='"Weekly,Every 2 weeks,Monthly,Quarterly,Twice a year,Yearly"', allow_blank=True)
st.add_data_validation(dv_b); st.add_data_validation(dv_o)
for rr in range(BR1, BR2 + 1):
    i = rr - BR1
    vals = bills[i] if i < len(bills) else (None, None, None, None, None)
    for col, v in zip((1, 2, 3, 4, 5), vals):
        c = st.cell(row=rr, column=col, value=v); inp(c, MONEY if col == 4 else None)
    dv_b.add(f"B{rr}"); dv_o.add(f"C{rr}")
    occ = f'IF(C{rr}="Weekly",52,IF(C{rr}="Every 2 weeks",26,IF(C{rr}="Monthly",12,IF(C{rr}="Quarterly",4,IF(C{rr}="Twice a year",2,IF(C{rr}="Yearly",1,0))))))'
    c = st.cell(row=rr, column=6, value=f'=IF(OR(D{rr}="",C{rr}=""),"",D{rr}*{occ})'); calc(c, MONEY)
    c = st.cell(row=rr, column=7, value=f'=IF(F{rr}="","",ROUNDUP(F{rr}/IF(OR(C{rr}="Weekly",C{rr}="Every 2 weeks"),$B$9,$B$10),0))'); calc(c, MONEY, bold=True)
    inp(st.cell(row=rr, column=8))
tr = BR2 + 1
st.cell(row=tr, column=1, value="Totals").font = font(bold=True)
c = st.cell(row=tr, column=6, value=f"=SUM(F{BR1}:F{BR2})"); calc(c, MONEY, bold=True)
c = st.cell(row=tr, column=7, value=f"=SUM(G{BR1}:G{BR2})"); calc(c, MONEY, bold=True)
st.cell(row=tr + 1, column=1, value="Per payday is rounded up to the next dollar, so a bucket is never short.").font = font(size=9, italic=True, color="6B6657")
st.freeze_panes = "A5"
st.sheet_properties.tabColor = GREEN

# ======================= PAYDAY =======================
p = wb.create_sheet("Payday")
title(p, "Payday", "Run this every payday. Same six steps, same order, every time.")
widths(p, {"A": 8, "B": 22, "C": 16, "D": 10, "E": 58})
p["B4"] = "Pay date"; p["B4"].font = font(bold=True); p["C4"] = "=Setup!B7"; inp(p["C4"], "ddd d mmm yyyy")
p["B5"] = "Paycheck this time"; p["B5"].font = font(bold=True); inp(p["C5"], MONEY)
p["D5"] = "Leave blank to use your usual paycheck."; p["D5"].font = font(size=9, italic=True, color="6B6657")
p["B6"] = "Using"; p["B6"].font = font(color="6B6657"); p["C6"] = '=IF(C5="",Setup!B5,C5)'; calc(p["C6"], MONEY, bold=True)
header(p, 8, ["Step", "Bucket", "Amount", "Done?", "How"])
how = ["Move it to your bill accounts, or pay what's due before next payday.",
       "Every minimum first. Any extra goes to the smallest balance.",
       "Move it now, before you can spend it.",
       "Move it into its own account. Out of sight.",
       "Groceries, gas, kids and fun, from what's left.",
       "What's left. Leave it alone until next payday."]
for i in range(6):
    rr = 9 + i; nm = names[i]
    c = p.cell(row=rr, column=1, value=i + 1); c.font = font(bold=True, color="FFFFFF"); c.fill = fill(BCOL[nm]); c.alignment = Alignment(horizontal="center")
    c = p.cell(row=rr, column=2, value=nm); c.font = font(bold=True, color=BCOL[nm]); c.fill = fill(TINT[nm])
    f_ = f"=Setup!B{15 + i}" if i < 5 else "=C6-SUM(C9:C13)"
    c = p.cell(row=rr, column=3, value=f_); calc(c, MONEY, bold=True); c.fill = fill(TINT[nm])
    c = p.cell(row=rr, column=4); inp(c); c.alignment = Alignment(horizontal="center")
    p.cell(row=rr, column=5, value=how[i]).font = font(size=9.5, color="5B6475")
    p.row_dimensions[rr].height = 22
dvp = DataValidation(type="list", formula1='"✓"', allow_blank=True); p.add_data_validation(dvp); dvp.add("D9:D14")
p["B16"] = "Safe to spend until next payday"; p["B16"].font = font(bold=True, color=GREEN, size=12)
p["C16"] = "=C13"; calc(p["C16"], MONEY, bold=True, color=GREEN)
p["B17"] = "Spare (buffer)"; p["B17"].font = font(bold=True)
p["C17"] = "=C14"; calc(p["C17"], MONEY, bold=True)
p["E17"] = '=IF(C14<Setup!B8,"⚠ Under your buffer target. Trim Everyday (step 5) first.","✓ Buffer is healthy.")'
p["E17"].font = font(bold=True, color=GREEN)
p.conditional_formatting.add("E17", FormulaRule(formula=["$C$14<Setup!$B$8"], font=Font(name=F, bold=True, color=RED)))
p["B19"] = '=IF(COUNTIF(D9:D14,"✓")=6,"All six done. Add a line to the Payday Log.",COUNTIF(D9:D14,"✓")&" of 6 done")'
p["B19"].font = font(bold=True, color=NAVY, size=11)
p["B21"] = "Pay the bills. Pay the debt. Pay yourself. Guard the rest."; p["B21"].font = font(bold=True, size=13, color=GREEN)
p["B22"] = "Clear the ticks (select D9:D14, press Delete) at the start of each payday."; p["B22"].font = font(size=9, italic=True, color="6B6657")
p.sheet_properties.tabColor = GOLD

# ======================= YEAR OF PAYDAYS =======================
y = wb.create_sheet("Year of Paydays")
title(y, "Year of Paydays", "Fills itself in from your next payday on the Setup tab. Extra paydays are starred.")
widths(y, {"A": 6, "B": 18, "C": 18, "D": 12, "E": 14, "F": 44})
y["A3"] = '="Extra paydays this year: "&COUNTIF(D6:D58,"★ EXTRA")'; y["A3"].font = font(bold=True, color=GOLD, size=11)
header(y, 5, ["#", "Payday", "Month", "Extra?", "Paycheck", "Big bills this month / where the extra goes"])
for n in range(1, 54):
    rr = 5 + n
    y.cell(row=rr, column=1, value=n).font = font(color="6B6657")
    first, fq = "Setup!$B$7", "Setup!$B$6"
    raw = (f'IF({fq}="Weekly",{first}+({n}-1)*7,IF({fq}="Every 2 weeks",{first}+({n}-1)*14,'
           f'IF({fq}="Twice a month",EDATE({first},INT(({n}-1)/2))+MOD({n}-1,2)*14,EDATE({first},{n}-1))))')
    c = y.cell(row=rr, column=2, value=f'=IF({first}="","",IF({raw}<EDATE({first},12),{raw},""))'); calc(c, "ddd d mmm yyyy")
    c = y.cell(row=rr, column=3, value=f'=IF(B{rr}="","",TEXT(B{rr},"mmmm yyyy"))'); calc(c)
    c = y.cell(row=rr, column=4, value=f'=IF(B{rr}="","",IF(COUNTIFS($B$6:B{rr},">="&DATE(YEAR(B{rr}),MONTH(B{rr}),1),$B$6:B{rr},"<="&B{rr})>Setup!$B$11,"★ EXTRA",""))'); calc(c, bold=True, color="8A6D12")
    c = y.cell(row=rr, column=5, value=f'=IF(B{rr}="","",Setup!$B$5)'); calc(c, MONEY)
    inp(y.cell(row=rr, column=6))
y.conditional_formatting.add("A6:F58", FormulaRule(formula=['$D6="★ EXTRA"'], fill=fill("FBF3D9")))
y.freeze_panes = "A6"
y.sheet_properties.tabColor = NAVY

# ======================= PAYDAY LOG =======================
lg = wb.create_sheet("Payday Log")
title(lg, "Payday Log", "One line per payday. Watch the minutes drop as it becomes routine.")
widths(lg, {"A": 6, "B": 16, "C": 13, "D": 14, "E": 13, "F": 12, "G": 10, "H": 11, "I": 15, "J": 36})
lg["A3"] = '="Paydays logged: "&COUNT(C7:C59)'; lg["A3"].font = font(bold=True, color=NAVY)
lg["D3"] = '=IF(COUNT(H7:H59)=0,"","Average minutes: "&ROUND(AVERAGE(H7:H59),0)&"   ·   Fastest: "&MIN(H7:H59))'; lg["D3"].font = font(bold=True, color=GREEN)
header(lg, 6, ["#", "Pay date", "Paycheck", "Into buckets 1–4", "Everyday", "Buffer", "All moved?", "Minutes to plan", "Buffer left before next payday", "Notes"])
dvl = DataValidation(type="list", formula1='"✓"', allow_blank=True); lg.add_data_validation(dvl); dvl.add("G7:G59")
for n in range(1, 54):
    rr = 6 + n
    lg.cell(row=rr, column=1, value=n).font = font(color="6B6657")
    c = lg.cell(row=rr, column=2, value=f"=IF('Year of Paydays'!B{5+n}=\"\",\"\",'Year of Paydays'!B{5+n})"); calc(c, "d mmm yyyy")
    inp(lg.cell(row=rr, column=3), MONEY)
    c = lg.cell(row=rr, column=4, value=f'=IF(C{rr}="","",SUM(Setup!$B$15:$B$18))'); calc(c, MONEY)
    c = lg.cell(row=rr, column=5, value=f'=IF(C{rr}="","",Setup!$B$19)'); calc(c, MONEY)
    c = lg.cell(row=rr, column=6, value=f'=IF(C{rr}="","",C{rr}-D{rr}-E{rr})'); calc(c, MONEY, bold=True)
    c = lg.cell(row=rr, column=7); inp(c); c.alignment = Alignment(horizontal="center")
    inp(lg.cell(row=rr, column=8), "0"); inp(lg.cell(row=rr, column=9), MONEY); inp(lg.cell(row=rr, column=10))
ch = LineChart(); ch.title = "Minutes to plan each payday"; ch.y_axis.title = "Minutes"; ch.x_axis.title = "Payday #"
ch.add_data(Reference(lg, min_col=8, min_row=6, max_row=30), titles_from_data=True)
ch.set_categories(Reference(lg, min_col=1, min_row=7, max_row=30))
ch.height, ch.width = 7, 16; ch.legend = None
lg.add_chart(ch, "L6")
lg.freeze_panes = "A7"
lg.sheet_properties.tabColor = GOLD

# ======================= 30-DAY CHALLENGE =======================
TASKS = ["Read pages 2 and 3", "Pull up 3 months of bank statements", "Highlight every payment that repeats", "Fill in the Bill Finder",
    "Hunt down 3 bills you forgot", "Do the Year of Paydays", "Rest day. Just look at your list.",
    "Read the 6 Buckets (page 5)", "Open or name your bucket accounts", "Fill in Bucket Setup", "Set up the automatic transfers",
    "Put your bills in the spreadsheet", "Show your partner the Payday Number", "Rest day",
    "Check: did every transfer go out?", "Buffer check. Still $100+?", "No-spend day", "Cancel one thing you don't use",
    "Shop from a list only", "Say the 6 buckets from memory", "Rest day",
    "Check every bucket account balance", "No-spend day", "Log this week's spending", "Move any leftover to Don't Touch",
    "Plan your next extra payday", "Do the monthly check-in", "Rest day",
    "Fill in your Payday Routine Card", "Celebrate. Something free, together."]
cc = wb.create_sheet("30-Day Challenge")
title(cc, "The 30-Day Lock-In Challenge", "One small thing a day. Missed a day? Don't restart. Just do today's.")
widths(cc, {"A": 8, "B": 16, "C": 46, "D": 10})
cc["B4"] = "Start date"; cc["B4"].font = font(bold=True); cc["C4"] = dt.date(2026, 10, 12); inp(cc["C4"], "ddd d mmm yyyy")
cc["B5"] = '="Days done: "&COUNTIF(D8:D37,"✓")&" / 30"'; cc["B5"].font = font(bold=True, size=12, color=GREEN)
cc["C5"] = '=REPT("■",COUNTIF(D8:D37,"✓"))&REPT("□",30-COUNTIF(D8:D37,"✓"))'; cc["C5"].font = font(color=GREEN, size=11)
header(cc, 7, ["Day", "Date", "Today's task", "Done?"])
dvc = DataValidation(type="list", formula1='"✓"', allow_blank=True); cc.add_data_validation(dvc); dvc.add("D8:D37")
for i, t in enumerate(TASKS):
    rr = 8 + i
    c = cc.cell(row=rr, column=1, value=i + 1); c.font = font(bold=True, color=NAVY); c.alignment = Alignment(horizontal="center")
    c = cc.cell(row=rr, column=2, value=f'=IF($C$4="","",$C$4+{i})'); calc(c, "ddd d mmm")
    c = cc.cell(row=rr, column=3, value=t); c.font = font(color="6B6657" if t.startswith("Rest") else "000000"); c.border = BOX
    c = cc.cell(row=rr, column=4); inp(c); c.alignment = Alignment(horizontal="center")
cc.conditional_formatting.add("A8:D37", FormulaRule(formula=['$D8="✓"'], fill=fill("E2F0E7")))
cc["C39"] = "Payday during the challenge? Do the Payday tab that day too."; cc["C39"].font = font(italic=True, color=NAVY)
cc.sheet_properties.tabColor = BLACK

# ======================= CHECK-IN =======================
ci = wb.create_sheet("Check-In")
title(ci, "The 10-Minute Monthly Check-In", "Once a month, sit down together. Ten minutes, six questions.")
widths(ci, {"A": 16, "B": 16, "C": 26, "D": 18, "E": 26, "F": 26, "G": 26})
header(ci, 4, ["Month", "Every transfer went out?", "A bill that wasn't in a bucket?", "Buffer left before paydays", "What surprised us?", "One thing we'll change", "One win to celebrate"])
dvy = DataValidation(type="list", formula1='"Yes,Mostly,No"', allow_blank=True); ci.add_data_validation(dvy); dvy.add("B5:B16")
for n in range(12):
    rr = 5 + n
    c = ci.cell(row=rr, column=1, value=f'=IF(Setup!$B$7="","",TEXT(EDATE(DATE(YEAR(Setup!$B$7),MONTH(Setup!$B$7),1),{n}),"mmmm yyyy"))'); calc(c, bold=True, color=NAVY)
    for col in range(2, 8):
        c = ci.cell(row=rr, column=col); inp(c); c.alignment = Alignment(wrap_text=True, vertical="top")
    ci.row_dimensions[rr].height = 30
ci.sheet_properties.tabColor = BLACK

# ======================= TRACKERS =======================
tk = wb.create_sheet("Trackers")
title(tk, "Watch It Grow", "Don't Touch fund, savings goal and debt snowball.")
widths(tk, {"A": 16, "B": 13, "C": 13, "D": 4, "E": 16, "F": 13, "G": 13, "H": 4, "I": 22, "J": 13, "K": 12, "L": 10, "M": 9})
def fund(col, name, color, goal, goal_label):
    L = lambda k: chr(ord(col) + k)
    tk[f"{L(0)}4"] = name; tk[f"{L(0)}4"].font = font(bold=True, size=11, color=color)
    tk[f"{L(0)}5"] = goal_label; tk[f"{L(0)}5"].font = font(bold=True)
    tk[f"{L(1)}5"] = goal; inp(tk[f"{L(1)}5"], MONEY)
    tk[f"{L(0)}6"] = "Balance"; tk[f"{L(0)}6"].font = font(bold=True)
    tk[f"{L(1)}6"] = f"=SUM({L(1)}10:{L(1)}60)-SUM({L(2)}10:{L(2)}60)"; calc(tk[f"{L(1)}6"], MONEY, bold=True, color=color)
    tk[f"{L(0)}7"] = f'=IF({L(1)}5>0,REPT("■",MIN(10,INT({L(1)}6/{L(1)}5*10)))&REPT("□",10-MIN(10,INT({L(1)}6/{L(1)}5*10)))&"  "&TEXT(MIN(1,{L(1)}6/{L(1)}5),"0%"),"")'
    tk[f"{L(0)}7"].font = font(bold=True, color=color, size=11)
    header(tk, 9, ["Date", "Added", "Used"], col=ord(col) - 64, bg=color)
    for rr in range(10, 61):
        inp(tk[f"{L(0)}{rr}"], "d mmm yyyy"); inp(tk[f"{L(1)}{rr}"], MONEY); inp(tk[f"{L(2)}{rr}"], MONEY)
fund("A", "DON'T TOUCH FUND", RED, 500, "Goal")
fund("E", "SAVINGS GOAL", GREEN, 1000, "Goal")
tk["E8"] = "Saving for:"; tk["E8"].font = font(bold=True); inp(tk["F8"]); tk.merge_cells("F8:G8")
tk["I4"] = "DEBT SNOWBALL"; tk["I4"].font = font(bold=True, size=11, color=BLACK)
tk["I5"] = "Total owed"; tk["I5"].font = font(bold=True); tk["J5"] = "=SUM(J10:J24)"; calc(tk["J5"], MONEY, bold=True)
tk["I6"] = "Total minimums"; tk["I6"].font = font(bold=True); tk["J6"] = "=SUM(K10:K24)"; calc(tk["J6"], MONEY, bold=True)
tk["I7"] = "Pay the smallest balance first. When it's gone, roll its payment into the next one."; tk["I7"].font = font(size=9, italic=True, color="6B6657")
header(tk, 9, ["Debt", "Balance", "Minimum", "Order", "Paid off"], col=9, bg=BLACK)
dvd = DataValidation(type="list", formula1='"✓"', allow_blank=True); tk.add_data_validation(dvd); dvd.add("M10:M24")
for rr in range(10, 25):
    inp(tk[f"I{rr}"]); inp(tk[f"J{rr}"], MONEY); inp(tk[f"K{rr}"], MONEY)
    tk[f"L{rr}"] = f'=IF(OR(J{rr}="",J{rr}=0),"",COUNTIF($J$10:$J$24,"<"&J{rr})-COUNTIF($J$10:$J$24,0)+1)'; calc(tk[f"L{rr}"], "0", bold=True)
    c = tk[f"M{rr}"]; inp(c); c.alignment = Alignment(horizontal="center")
tk.conditional_formatting.add("I10:M24", FormulaRule(formula=['$M10="✓"'], fill=fill("E2F0E7")))
tk.sheet_properties.tabColor = RED

for ws in wb.worksheets:
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT)
print("saved", OUT)
