import sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.utils import get_column_letter as CL

OUT = sys.argv[1] if len(sys.argv) > 1 else "The Paycheck Budget System - Spreadsheet.xlsx"
NAVY, GREEN, GOLD, RED, BROWN, BLACK = "1B3A6B", "1E7A4C", "C9A227", "B23A3A", "7A5A0F", "121212"
INPUT = "FFF4CC"; GREYT = "6B6657"
TINT = {"Bills": "E6ECF5", "Debt": "ECEAE4", "Savings": "E2F0E7", "Don't Touch": "F7E6E6", "Everyday": "F6EED3", "Buffer": "FBF3D9"}
BCOL = {"Bills": NAVY, "Debt": BLACK, "Savings": GREEN, "Don't Touch": RED, "Everyday": BROWN, "Buffer": "8A6D12"}
F = "Arial"
MONEY = '$#,##0;-$#,##0;"-"'
MONEY_BLANK = '$#,##0;-$#,##0;""'
DATE = "ddd d mmm yyyy"
thin = Side(style="thin", color="D9D3C4"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

def font(**k):
    k.setdefault("name", F); k.setdefault("size", 10); return Font(**k)
def fill(c): return PatternFill("solid", start_color=c, end_color=c)
def title(ws, t, sub):
    ws["A1"] = t; ws["A1"].font = font(size=20, bold=True, color=NAVY); ws.row_dimensions[1].height = 30
    ws["A2"] = sub; ws["A2"].font = font(size=11, color=GREEN, bold=True); ws.sheet_view.showGridLines = False
def header(ws, row, labels, col=1, bg=NAVY, fg="FFFFFF", h=30):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=l)
        c.font = font(bold=True, color=fg, size=9.5); c.fill = fill(bg); c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = h
def step(ws, ref, text, color=NAVY):
    ws[ref] = text; ws[ref].font = font(bold=True, size=12, color=color)
def inp(c, fmt=None):
    c.fill = fill(INPUT); c.font = font(color="0000FF"); c.border = BOX
    if fmt: c.number_format = fmt
def calc(c, fmt=None, bold=False, color="000000"):
    c.font = font(bold=bold, color=color); c.border = BOX
    if fmt: c.number_format = fmt
def widths(ws, d):
    for k, v in d.items(): ws.column_dimensions[k].width = v

wb = Workbook()
B = "'My Budget'!"

# ======================= START HERE =======================
s = wb.active; s.title = "Start Here"
title(s, "The Paycheck Budget System", "Companion spreadsheet · The Payday Plan")
widths(s, {"A": 5, "B": 100})
lines = [
    ("H", "IT'S THREE STEPS"),
    (1, "Open the My Budget tab and put your pay in: how much, how often, and the date of your next payday."),
    (2, "Put your bills in, with how often they come and when the next one is due. Then your everyday spending."),
    (3, "That's it. The sheet works out the rest."),
    ("", ""),
    ("H", "WHAT IT WORKS OUT FOR YOU"),
    ("•", "My Budget: how much goes into each bucket every payday, and your Payday Number."),
    ("•", "Pay Calendar: every payday for the next 12 months, the extra paydays, and exactly which bills land in which pay."),
    ("•", "Payday: the six steps to tick off each payday."),
    ("", ""),
    ("H", "GOOD TO KNOW"),
    ("", "Yellow boxes with blue text are yours to type in. White boxes work themselves out, so leave them alone."),
    ("", "The example numbers are a family paid $2,500 every two weeks. Type over them with your own."),
    ("", "A bill that's due soon may not have had time to build up yet. The sheet tells you how much to top it up by, once. After that, the buckets keep up on their own."),
    ("", "Works in Excel and Google Sheets (upload to Google Drive, then open with Google Sheets)."),
    ("", ""),
    ("", "Pay the bills. Pay the debt. Pay yourself. Guard the rest."),
    ("", "General information, not financial advice. For personal use only. © The Payday Plan"),
]
r = 4
for a, b in lines:
    if a == "H":
        s.cell(row=r, column=1, value=b).font = font(bold=True, size=11, color=NAVY)
    else:
        if isinstance(a, int) or a == "•":
            c = s.cell(row=r, column=1, value=a); c.font = font(bold=True, color="FFFFFF" if isinstance(a, int) else NAVY)
            if isinstance(a, int): c.fill = fill(NAVY)
            c.alignment = Alignment(horizontal="center", vertical="top")
        c = s.cell(row=r, column=2, value=b or None); c.font = font(size=10.5); c.alignment = Alignment(wrap_text=True, vertical="top")
        if b.startswith("Yellow"): c.fill = fill(INPUT); c.font = font(size=10.5, color="0000FF")
        if b.startswith("Pay the bills"): c.font = font(size=13, bold=True, color=GREEN)
        if b.startswith("General"): c.font = font(size=9, color=GREYT)
        if len(b) > 105: s.row_dimensions[r].height = 28
    r += 1
s.sheet_properties.tabColor = NAVY

# ======================= MY BUDGET =======================
m = wb.create_sheet("My Budget")
title(m, "My Budget", "Fill in the yellow boxes. Everything else works itself out.")
widths(m, {"A": 26, "B": 14, "C": 15, "D": 12, "E": 16, "F": 12, "G": 15, "H": 20, "I": 3, "J": 18, "K": 14, "L": 14, "M": 18})
step(m, "A4", "STEP 1 · YOUR PAY")
pay = [("Paycheck amount", 2500, MONEY), ("How often you're paid", "Every 2 weeks", None), ("Your next payday", dt.date(2026, 10, 16), DATE),
       ("Buffer to keep each payday", 100, MONEY)]
for i, (lab, v, fm) in enumerate(pay):
    rr = 5 + i
    m.cell(row=rr, column=1, value=lab).font = font(bold=True)
    c = m.cell(row=rr, column=2, value=v); inp(c, fm)
dv_pay = DataValidation(type="list", formula1='"Weekly,Every 2 weeks,Twice a month,Monthly"'); m.add_data_validation(dv_pay); dv_pay.add("B6")
# hidden helpers (S:T)
m["S5"] = "Paydays a year"; m["T5"] = '=IF(B6="Weekly",52,IF(B6="Every 2 weeks",26,IF(B6="Twice a month",24,12)))'
m["S6"] = "Divide-by"; m["T6"] = '=IF(B6="Weekly",48,IF(B6="Every 2 weeks",24,IF(B6="Twice a month",24,12)))'
m["S7"] = "Usual paydays a month"; m["T7"] = '=IF(B6="Weekly",4,IF(B6="Monthly",1,2))'

R1, R2 = 14, 43
step(m, "A11", "STEP 2 · YOUR BILLS")
m["A12"] = "Everything that isn't day-to-day spending: bills, debt repayments, savings and set-asides. The next due date lets the Pay Calendar show when each one lands."
m["A12"].font = font(size=9.5, color=GREYT); m["A12"].alignment = Alignment(wrap_text=True, vertical="top"); m.merge_cells("A12:H12"); m.row_dimensions[12].height = 26
header(m, 13, ["Bill", "Bucket", "How often", "Amount", "Next due date", "Per payday", "Ready by the due date?", "Notes"])
bills = [
    ("Rent", "Bills", "Monthly", 1800, dt.date(2026, 11, 1)), ("Power", "Bills", "Monthly", 160, dt.date(2026, 11, 20)),
    ("Phone & internet", "Bills", "Monthly", 140, dt.date(2026, 11, 5)), ("Car insurance", "Bills", "Yearly", 1440, dt.date(2027, 8, 15)),
    ("Car registration", "Bills", "Yearly", 360, dt.date(2027, 3, 1)), ("Water", "Bills", "Quarterly", 150, dt.date(2026, 12, 15)),
    ("Car service & tires", "Bills", "Yearly", 600, dt.date(2027, 6, 1)), ("Streaming & apps", "Bills", "Monthly", 50, dt.date(2026, 11, 10)),
    ("Credit card minimum", "Debt", "Monthly", 200, dt.date(2026, 11, 12)),
    ("Savings", "Savings", "Monthly", 160, None), ("Christmas", "Savings", "Yearly", 720, dt.date(2026, 12, 10)),
    ("Birthdays & gifts", "Savings", "Yearly", 480, None), ("School costs", "Savings", "Yearly", 480, dt.date(2027, 1, 25)),
    ("Don't Touch fund", "Don't Touch", "Monthly", 160, None),
]
dv_b = DataValidation(type="list", formula1='"Bills,Debt,Savings,Don\'t Touch"', allow_blank=True)
dv_o = DataValidation(type="list", formula1='"Weekly,Every 2 weeks,Monthly,Quarterly,Twice a year,Yearly"', allow_blank=True)
m.add_data_validation(dv_b); m.add_data_validation(dv_o)
for rr in range(R1, R2 + 1):
    i = rr - R1
    vals = bills[i] if i < len(bills) else (None,) * 5
    for col, v in zip(range(1, 6), vals):
        c = m.cell(row=rr, column=col, value=v); inp(c, MONEY if col == 4 else (DATE if col == 5 else None))
    dv_b.add(f"B{rr}"); dv_o.add(f"C{rr}")
    occ = f'IF(C{rr}="Weekly",52,IF(C{rr}="Every 2 weeks",26,IF(C{rr}="Monthly",12,IF(C{rr}="Quarterly",4,IF(C{rr}="Twice a year",2,IF(C{rr}="Yearly",1,0))))))'
    c = m.cell(row=rr, column=6, value=f'=IF(OR(D{rr}="",C{rr}=""),"",ROUNDUP(D{rr}*{occ}/IF(OR(C{rr}="Weekly",C{rr}="Every 2 weeks"),$T$5,$T$6),0))'); calc(c, MONEY, bold=True)
    c = m.cell(row=rr, column=7, value=f'=IF(OR(D{rr}="",C{rr}="",E{rr}=""),"",MAX(0,D{rr}-F{rr}*COUNTIF(\'Pay Calendar\'!$B$8:$B$61,"<="&E{rr})))')
    calc(c, '"Top up "$#,##0;;"✓ Ready"')
    inp(m.cell(row=rr, column=8))
    m[f"P{rr}"] = f'=IF(C{rr}="Weekly",7,IF(C{rr}="Every 2 weeks",14,0))'
    m[f"Q{rr}"] = f'=IF(C{rr}="Monthly",1,IF(C{rr}="Quarterly",3,IF(C{rr}="Twice a year",6,IF(C{rr}="Yearly",12,0))))'
m.conditional_formatting.add(f"G{R1}:G{R2}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=F, bold=True, color=RED)))
m.conditional_formatting.add(f"G{R1}:G{R2}", CellIsRule(operator="equal", formula=["0"], font=Font(name=F, color=GREEN)))
TR = R2 + 1
m.cell(row=TR, column=1, value="Totals").font = font(bold=True)
c = m.cell(row=TR, column=6, value=f"=SUM(F{R1}:F{R2})"); calc(c, MONEY, bold=True)
c = m.cell(row=TR, column=7, value=f"=SUM(G{R1}:G{R2})"); calc(c, MONEY, bold=True, color=RED)

E1, E2 = 49, 56
step(m, "A47", "STEP 3 · YOUR EVERYDAY SPENDING", BROWN)
header(m, 48, ["Item", "Each payday"], bg=BROWN, h=22)
every = [("Groceries", 450), ("Gas", 150), ("Kids & household", 150), ("Fun money", 120)]
for rr in range(E1, E2 + 1):
    i = rr - E1
    a, v = every[i] if i < len(every) else (None, None)
    inp(m.cell(row=rr, column=1, value=a)); inp(m.cell(row=rr, column=2, value=v), MONEY)
m.cell(row=E2 + 1, column=1, value="Everyday total").font = font(bold=True, color=BROWN)
c = m.cell(row=E2 + 1, column=2, value=f"=SUM(B{E1}:B{E2})"); calc(c, MONEY, bold=True)
ETOT = f"B{E2 + 1}"

# results box (J:M)
m["J4"] = "YOUR PAYDAY NUMBER"; m["J4"].font = font(bold=True, size=12, color="FFFFFF"); m["J4"].fill = fill(NAVY)
for col in "KLM": m[f"{col}4"].fill = fill(NAVY)
names = ["Bills", "Debt", "Savings", "Don't Touch", "Everyday", "Buffer"]
for i, nm in enumerate(names):
    rr = 5 + i
    c = m.cell(row=rr, column=10, value=f"{i+1}  {nm}"); c.font = font(bold=True, color=BCOL[nm]); c.fill = fill(TINT[nm]); c.border = BOX
    if i < 4:
        f_ = f'=SUMIF($B${R1}:$B${R2},"{nm.replace(chr(39), chr(39))}",$F${R1}:$F${R2})'
    elif i == 4:
        f_ = f"={ETOT}"
    else:
        f_ = "=B5-SUM(K5:K9)"
    c = m.cell(row=rr, column=11, value=f_); calc(c, MONEY, bold=True); c.fill = fill(TINT[nm])
m["J12"] = '=IF(K10>=B8,"✓ Your buffer is healthy.","⚠ Buffer is under "&TEXT(B8,"$#,##0")&". Trim Everyday first.")'
m["J12"].font = font(bold=True, color=GREEN)
m.conditional_formatting.add("J12", FormulaRule(formula=["$K$10<$B$8"], font=Font(name=F, bold=True, color=RED)))
m.conditional_formatting.add("K10", CellIsRule(operator="lessThan", formula=["$B$8"], fill=fill("F7E6E6"), font=Font(name=F, bold=True, color=RED)))
sent = [('="Every payday, "&TEXT(SUM(K5:K8),"$#,##0")&" moves into buckets 1 to 4."', 14),
        ('=TEXT(K9,"$#,##0")&" is yours to live on until next payday."', 15),
        ('="And "&TEXT(K10,"$#,##0")&" stays spare."', 16)]
for f_, rr in sent:
    m[f"J{rr}"] = f_; m[f"J{rr}"].font = font(size=11.5, bold=True, color=NAVY); m.merge_cells(f"J{rr}:M{rr}")
m["J18"] = f'=IF(G{TR}>0,"One-off catch-up: "&TEXT(G{TR},"$#,##0"),"✓ Every bill is ready by its due date.")'
m["J18"].font = font(size=11.5, bold=True, color=RED); m.merge_cells("J18:M18")
m["J19"] = f'=IF(G{TR}>0,"A few bills are due before their bucket has filled.","")'
m["J20"] = f'=IF(G{TR}>0,"Top them up once (column G). After that, they keep up.","")'
for a in ("J19", "J20"):
    m[a].font = font(size=10, color=RED); m.merge_cells(f"{a}:M{a[1:]}")
m.conditional_formatting.add("J18", FormulaRule(formula=[f"$G${TR}=0"], font=Font(name=F, bold=True, color=GREEN)))
for col in "PQST": m.column_dimensions[col].hidden = True
m.freeze_panes = "A4"
m.sheet_properties.tabColor = GREEN

# ======================= PAY CALENDAR =======================
pc = wb.create_sheet("Pay Calendar")
title(pc, "Pay Calendar", "Every payday for the next 12 months, and which bills land in each pay.")
NB = R2 - R1 + 1  # 30 bill columns
G0 = 9  # first bill column (I)
widths(pc, {"A": 5, "B": 16, "C": 8, "D": 11, "E": 12, "F": 11, "G": 10, "H": 12})
for j in range(NB): pc.column_dimensions[CL(G0 + j)].width = 11
ANCH, FREQ = f"{B}$B$7", f"{B}$B$6"
def raw(n):
    if n % 2 == 1:
        twice = f"EDATE({ANCH},{(n - 1) // 2})"
    else:
        a = f"DAY({ANCH})"
        x = f"EDATE(DATE(YEAR({ANCH}),MONTH({ANCH}),1),IF({a}<=15,{(n - 2) // 2},{n // 2}))"
        twice = f"({x}+MIN(IF({a}<=15,{a}+15,{a}-15),DAY(EOMONTH({x},0)))-1)"
    return (f'IF({FREQ}="Weekly",{ANCH}+({n}-1)*7,IF({FREQ}="Every 2 weeks",{ANCH}+({n}-1)*14,'
            f'IF({FREQ}="Twice a month",{twice},EDATE({ANCH},{n}-1))))')
FIRST, LAST = 8, 61  # 54 rows
HELP = CL(G0 + NB + 1)  # hidden 'period ends' column
pc["A3"] = f'="Paydays: "&COUNT(B{FIRST}:B{LAST})&"   ·   Extra paydays: "&COUNTIF(C{FIRST}:C{LAST},"★")'
pc["A3"].font = font(bold=True, color=NAVY, size=11)
pc["A4"] = f'=IF(MAX(H{FIRST}:H{LAST})=0,"","Biggest pay for bills: "&TEXT(MAX(H{FIRST}:H{LAST}),"$#,##0")&" in the pay starting "&TEXT(INDEX(B{FIRST}:B{LAST},MATCH(MAX(H{FIRST}:H{LAST}),H{FIRST}:H{LAST},0)),"d mmm yyyy")&". Your buckets have it covered.")'
pc["A4"].font = font(bold=True, color=GREEN)
pc["A5"] = "Gold = a pile-up pay: more bills land than you set aside that payday. That's normal. The buckets cover it."
pc["A5"].font = font(size=9, italic=True, color=GREYT)
header(pc, 7, ["#", "Payday", "Extra?", "Pay", "Into buckets 1–4", "Everyday", "Buffer", "Bills due this pay"], h=40)
for j in range(NB):
    c = pc.cell(row=7, column=G0 + j, value=f'=IF({B}$A${R1 + j}="","",{B}$A${R1 + j})')
    c.font = font(bold=True, color="FFFFFF", size=9); c.fill = fill("2D4F86"); c.alignment = Alignment(wrap_text=True, vertical="center")
for n in range(1, LAST - FIRST + 2):
    rr = FIRST + n - 1
    pc.cell(row=rr, column=1, value=n).font = font(color=GREYT)
    c = pc.cell(row=rr, column=2, value=f'=IF({ANCH}="","",IF({raw(n)}<EDATE({ANCH},12),{raw(n)},""))'); calc(c, DATE)
    pc[f"{HELP}{rr}"] = f'=IF({ANCH}="","",{raw(n + 1)})'
    c = pc.cell(row=rr, column=3, value=f'=IF(B{rr}="","",IF(COUNTIFS($B${FIRST}:B{rr},">="&DATE(YEAR(B{rr}),MONTH(B{rr}),1),$B${FIRST}:B{rr},"<="&B{rr})>{B}$T$7,"★",""))')
    calc(c, bold=True, color="8A6D12"); c.alignment = Alignment(horizontal="center")
    c = pc.cell(row=rr, column=4, value=f'=IF(B{rr}="","",{B}$B$5)'); calc(c, MONEY)
    c = pc.cell(row=rr, column=5, value=f'=IF(B{rr}="","",SUM({B}$K$5:$K$8))'); calc(c, MONEY)
    c = pc.cell(row=rr, column=6, value=f'=IF(B{rr}="","",{B}$K$9)'); calc(c, MONEY)
    c = pc.cell(row=rr, column=7, value=f'=IF(B{rr}="","",{B}$K$10)'); calc(c, MONEY)
    c = pc.cell(row=rr, column=8, value=f'=IF(B{rr}="","",SUM({CL(G0)}{rr}:{CL(G0 + NB - 1)}{rr}))'); calc(c, MONEY, bold=True)
    s_, e_ = f"$B{rr}", f"${HELP}{rr}"
    for j in range(NB):
        R = R1 + j
        A_, C_, D_, P_, Q_ = f"{B}$D${R}", f"{B}$C${R}", f"{B}$E${R}", f"{B}$P${R}", f"{B}$Q${R}"
        f0 = f"IF({D_}>={s_},{D_},{D_}-INT(({D_}-{s_})/{P_})*{P_})"
        daycnt = f"IF({f0}<{e_},INT(({e_}-1-{f0})/{P_})+1,0)"
        q = f"INT(((YEAR({s_})-YEAR({D_}))*12+MONTH({s_})-MONTH({D_}))/{Q_})"
        mon = "+".join(f"(EDATE({D_},({q}+{k})*{Q_})>={s_})*(EDATE({D_},({q}+{k})*{Q_})<{e_})*(({q}+{k})>=0)" for k in range(3))
        f_ = f'=IF(OR({s_}="",{A_}="",{C_}="",{D_}=""),"",{A_}*IF({P_}>0,{daycnt},{mon}))'
        c = pc.cell(row=rr, column=G0 + j, value=f_); c.number_format = MONEY_BLANK; c.font = font(size=9.5); c.border = BOX
pc.conditional_formatting.add(f"{CL(G0)}{FIRST}:{CL(G0 + NB - 1)}{LAST}", CellIsRule(operator="greaterThan", formula=["0"], fill=fill("FBF3D9"), font=Font(name=F, bold=True, color=NAVY)))
pc.conditional_formatting.add(f"H{FIRST}:H{LAST}", FormulaRule(formula=[f'AND($B{FIRST}<>"",$H{FIRST}>$E{FIRST})'], fill=fill("F6E3A8")))
pc.conditional_formatting.add(f"A{FIRST}:I{LAST}", FormulaRule(formula=[f'$C{FIRST}="★"'], fill=fill("FBF3D9")))
pc["A5"] = "Gold = a pile-up pay: more bills land than you set aside that payday. That's normal, it's what the buckets are for. ★ = an extra payday."
pc.column_dimensions[HELP].hidden = True
pc.freeze_panes = pc[f"C{FIRST}"]
pc.sheet_properties.tabColor = NAVY

# ======================= PAYDAY =======================
p = wb.create_sheet("Payday")
title(p, "Payday", "Six steps, same order, every payday.")
widths(p, {"A": 7, "B": 18, "C": 14, "D": 10, "E": 60})
p["B4"] = "Paycheck this time"; p["B4"].font = font(bold=True); inp(p["C4"], MONEY)
p["D4"] = "Leave blank if it's your usual pay."; p["D4"].font = font(size=9, italic=True, color=GREYT)
p["B5"] = "Using"; p["B5"].font = font(color=GREYT); p["C5"] = f'=IF(C4="",{B}B5,C4)'; calc(p["C5"], MONEY, bold=True)
header(p, 7, ["Step", "Bucket", "Amount", "Done?", "How"])
how = ["Move it to your bill accounts, or pay what's due before next payday.", "Every minimum first. Any extra goes to the smallest balance.",
       "Move it now, before you can spend it.", "Move it into its own account. Out of sight.",
       "Groceries, gas, kids and fun, from what's left.", "What's left. Leave it alone until next payday."]
for i, nm in enumerate(names):
    rr = 8 + i
    c = p.cell(row=rr, column=1, value=i + 1); c.font = font(bold=True, color="FFFFFF"); c.fill = fill(BCOL[nm]); c.alignment = Alignment(horizontal="center")
    c = p.cell(row=rr, column=2, value=nm); c.font = font(bold=True, color=BCOL[nm]); c.fill = fill(TINT[nm])
    c = p.cell(row=rr, column=3, value=f"={B}K{5 + i}" if i < 5 else "=C5-SUM(C8:C12)"); calc(c, MONEY, bold=True); c.fill = fill(TINT[nm])
    c = p.cell(row=rr, column=4); inp(c); c.alignment = Alignment(horizontal="center")
    p.cell(row=rr, column=5, value=how[i]).font = font(size=9.5, color="5B6475"); p.row_dimensions[rr].height = 22
dvp = DataValidation(type="list", formula1='"✓"', allow_blank=True); p.add_data_validation(dvp); dvp.add("D8:D13")
p["B15"] = '=IF(COUNTIF(D8:D13,"✓")=6,"✓ All six done. See you next payday.",COUNTIF(D8:D13,"✓")&" of 6 done")'; p["B15"].font = font(bold=True, color=NAVY, size=11)
p["B16"] = f'=IF(C13<{B}B8,"⚠ Buffer is under your target. Trim Everyday (step 5) first.","")'; p["B16"].font = font(bold=True, color=RED)
p["B18"] = "Pay the bills. Pay the debt. Pay yourself. Guard the rest."; p["B18"].font = font(bold=True, size=13, color=GREEN)
p["B19"] = "Next payday: select the ticks in column D and press Delete."; p["B19"].font = font(size=9, italic=True, color=GREYT)
p.sheet_properties.tabColor = GOLD

# ======================= 30-DAY CHALLENGE =======================
TASKS = ["Read pages 2 and 3", "Pull up 3 months of bank statements", "Highlight every payment that repeats", "Fill in the Bill Finder",
    "Hunt down 3 bills you forgot", "Do the Year of Paydays", "Rest day. Just look at your list.",
    "Read the 6 Buckets (page 5)", "Open or name your bucket accounts", "Fill in Bucket Setup", "Set up the automatic transfers",
    "Put your bills in the spreadsheet", "Show your partner the Payday Number", "Rest day",
    "Check: did every transfer go out?", "Buffer check. Still $100+?", "No-spend day", "Cancel one thing you don't use",
    "Shop from a list only", "Say the 6 buckets from memory", "Rest day",
    "Check every bucket account balance", "No-spend day", "Look at this week's spending", "Move any leftover to Don't Touch",
    "Plan your next extra payday", "Do the monthly check-in", "Rest day",
    "Fill in your Payday Routine Card", "Celebrate. Something free, together."]
cc = wb.create_sheet("30-Day Challenge")
title(cc, "The 30-Day Lock-In Challenge", "One small thing a day. Missed a day? Don't restart. Just do today's.")
widths(cc, {"A": 8, "B": 16, "C": 46, "D": 10})
cc["B4"] = "Start date"; cc["B4"].font = font(bold=True); cc["C4"] = dt.date(2026, 10, 12); inp(cc["C4"], DATE)
cc["B5"] = '="Days done: "&COUNTIF(D8:D37,"✓")&" / 30"'; cc["B5"].font = font(bold=True, size=12, color=GREEN)
header(cc, 7, ["Day", "Date", "Today's task", "Done?"], h=22)
dvc = DataValidation(type="list", formula1='"✓"', allow_blank=True); cc.add_data_validation(dvc); dvc.add("D8:D37")
for i, t in enumerate(TASKS):
    rr = 8 + i
    c = cc.cell(row=rr, column=1, value=i + 1); c.font = font(bold=True, color=NAVY); c.alignment = Alignment(horizontal="center")
    c = cc.cell(row=rr, column=2, value=f'=IF($C$4="","",$C$4+{i})'); calc(c, "ddd d mmm")
    c = cc.cell(row=rr, column=3, value=t); c.font = font(color=GREYT if t.startswith("Rest") else "000000"); c.border = BOX
    c = cc.cell(row=rr, column=4); inp(c); c.alignment = Alignment(horizontal="center")
cc.conditional_formatting.add("A8:D37", FormulaRule(formula=['$D8="✓"'], fill=fill("E2F0E7")))
cc.sheet_properties.tabColor = BLACK

# ======================= CHECK-IN =======================
ci = wb.create_sheet("Check-In")
title(ci, "The 10-Minute Monthly Check-In", "Once a month, sit down together. Ten minutes, six questions.")
widths(ci, {"A": 16, "B": 16, "C": 26, "D": 18, "E": 26, "F": 26, "G": 26})
header(ci, 4, ["Month", "Every transfer went out?", "A bill that wasn't in a bucket?", "Buffer left before paydays", "What surprised us?", "One thing we'll change", "One win to celebrate"])
dvy = DataValidation(type="list", formula1='"Yes,Mostly,No"', allow_blank=True); ci.add_data_validation(dvy); dvy.add("B5:B16")
for n in range(12):
    rr = 5 + n
    c = ci.cell(row=rr, column=1, value=f'=IF({ANCH}="","",TEXT(EDATE(DATE(YEAR({ANCH}),MONTH({ANCH}),1),{n}),"mmmm yyyy"))'); calc(c, bold=True, color=NAVY)
    for col in range(2, 8):
        c = ci.cell(row=rr, column=col); inp(c); c.alignment = Alignment(wrap_text=True, vertical="top")
    ci.row_dimensions[rr].height = 30
ci.sheet_properties.tabColor = BLACK

# ======================= TRACKERS =======================
tk = wb.create_sheet("Trackers")
title(tk, "Watch It Grow", "Don't Touch fund, savings goal and debt snowball.")
widths(tk, {"A": 16, "B": 13, "C": 13, "D": 4, "E": 16, "F": 13, "G": 13, "H": 4, "I": 22, "J": 13, "K": 12, "L": 10, "M": 9})
def fund(col, name, color, goal):
    L = lambda k: chr(ord(col) + k)
    tk[f"{L(0)}4"] = name; tk[f"{L(0)}4"].font = font(bold=True, size=11, color=color)
    tk[f"{L(0)}5"] = "Goal"; tk[f"{L(0)}5"].font = font(bold=True); tk[f"{L(1)}5"] = goal; inp(tk[f"{L(1)}5"], MONEY)
    tk[f"{L(0)}6"] = "Balance"; tk[f"{L(0)}6"].font = font(bold=True)
    tk[f"{L(1)}6"] = f"=SUM({L(1)}10:{L(1)}60)-SUM({L(2)}10:{L(2)}60)"; calc(tk[f"{L(1)}6"], MONEY, bold=True, color=color)
    tk[f"{L(0)}7"] = f'=IF({L(1)}5>0,REPT("■",MAX(0,MIN(10,INT({L(1)}6/{L(1)}5*10))))&REPT("□",10-MAX(0,MIN(10,INT({L(1)}6/{L(1)}5*10))))&"  "&TEXT(MAX(0,MIN(1,{L(1)}6/{L(1)}5)),"0%"),"")'
    tk[f"{L(0)}7"].font = font(bold=True, color=color, size=11)
    header(tk, 9, ["Date", "Added", "Used"], col=ord(col) - 64, bg=color, h=20)
    for rr in range(10, 61):
        inp(tk[f"{L(0)}{rr}"], "d mmm yyyy"); inp(tk[f"{L(1)}{rr}"], MONEY); inp(tk[f"{L(2)}{rr}"], MONEY)
fund("A", "DON'T TOUCH FUND", RED, 500)
fund("E", "SAVINGS GOAL", GREEN, 1000)
tk["E8"] = "Saving for:"; tk["E8"].font = font(bold=True); inp(tk["F8"]); tk.merge_cells("F8:G8")
tk["I4"] = "DEBT SNOWBALL"; tk["I4"].font = font(bold=True, size=11, color=BLACK)
tk["I5"] = "Total owed"; tk["I5"].font = font(bold=True); tk["J5"] = "=SUM(J10:J24)"; calc(tk["J5"], MONEY, bold=True)
tk["I6"] = "Total minimums"; tk["I6"].font = font(bold=True); tk["J6"] = "=SUM(K10:K24)"; calc(tk["J6"], MONEY, bold=True)
tk["I7"] = "Pay the smallest balance first. When it's gone, roll its payment into the next one."; tk["I7"].font = font(size=9, italic=True, color=GREYT)
header(tk, 9, ["Debt", "Balance", "Minimum", "Order", "Paid off"], col=9, bg=BLACK, h=20)
dvd = DataValidation(type="list", formula1='"✓"', allow_blank=True); tk.add_data_validation(dvd); dvd.add("M10:M24")
for rr in range(10, 25):
    inp(tk[f"I{rr}"]); inp(tk[f"J{rr}"], MONEY); inp(tk[f"K{rr}"], MONEY)
    tk[f"L{rr}"] = f'=IF(OR(J{rr}="",J{rr}=0),"",COUNTIF($J$10:$J$24,"<"&J{rr})-COUNTIF($J$10:$J$24,0)+1)'; calc(tk[f"L{rr}"], "0", bold=True)
    c = tk[f"M{rr}"]; inp(c); c.alignment = Alignment(horizontal="center")
tk.conditional_formatting.add("I10:M24", FormulaRule(formula=['$M10="✓"'], fill=fill("E2F0E7")))
tk.sheet_properties.tabColor = RED

for ws in wb.worksheets:
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT)
print("saved", OUT)
